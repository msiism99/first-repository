"""Physical rotation/subspace checks and an actual-kernel, saved-file integration test."""
import ast
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

import nbformat
from nbclient import NotebookClient
import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
BIG_PHOTO = "9007199254740993"


def fixture_frame():
    rows = []
    for wafer in [1, 2, 3, 99]:
        for shot_x in [0., 140000.]:
            for x in np.linspace(-12000., 12000., 9):
                for y in np.linspace(-20000., 20000., 9):
                    if wafer == 1:
                        error = .004*(x/12000.)**4
                    elif wafer == 2:
                        error = .004*(y/20000.)**3
                    else:
                        error = .001*(x/12000.)**3+.002*(y/20000.)**2
                    rows.append(dict(LOT_ID="LOT_ROT", Wafer=wafer, photo_transn_seq=BIG_PHOTO, fcp_x=shot_x, fcp_y=0., coordinate_X=x, coordinate_Y=y, raw_minus_total_fit_x=error, raw_x=error+9., X_reg=error+99., STEP_PITCH_X=26000., STEP_PITCH_Y=44000., CHIP_X_NUM=2, CHIP_Y_NUM=2, TEST=len(rows)))
    rows.append({**rows[0], "raw_minus_total_fit_x": np.nan, "TEST": len(rows)})
    return pd.DataFrame(rows)


def fixture_files(base):
    stages = base / "output_decorrect_notebook/stages"
    stages.mkdir(parents=True)
    data = fixture_frame()
    data.to_parquet(stages / "21_OCO_WK_RK_REMOVED.parquet", index=False)
    (stages / "21_ADI_WK_RK_REMOVED.parquet").write_bytes(b"bad parquet - full Excel fallback required")
    data.to_excel(stages / "21_ADI_WK_RK_REMOVED.xlsx", index=False)
    data.head(1).to_excel(stages / "21_ADI_WK_RK_REMOVED_sample_1.xlsx", index=False)
    context = base / "output_wafer_context"
    context.mkdir()
    pd.DataFrame([dict(LOT_ID="LOT_ROT", Wafer=wafer, photo_transn_seq=BIG_PHOTO, product_code="ROTATED" if wafer != 99 else "OTHER", ReticleID="RET_A") for wafer in [1,2,3,99]]).to_parquet(context / "wafer_context.parquet", index=False)
    (stages.parent / "extension_settings.json").write_text(json.dumps(dict(cpe_options=dict(gopt_type=1, gopt_spec=.8, span_mode="OVO3_38P_F2F"), virtual_wafer_radius_um=144000., use_virtual_chip_points=True)), encoding="utf-8")


def modified_notebook(base):
    notebook = nbformat.read(HERE / "CPE_Critical_Rotation.ipynb", as_version=4)
    nbformat.validate(notebook)
    for cell in notebook.cells:
        if "rotation-00" not in cell.metadata.get("tags", []):
            continue
        replacements = dict(BASE_DIR=f"Path({str(base)!r})", FILTERS='{"product_code": ["ROTATED"]}', SHOW_PLOTS="False")
        lines = cell.source.splitlines()
        for node in ast.parse(cell.source).body:
            if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name) and node.targets[0].id in replacements:
                key = node.targets[0].id
                lines[node.lineno-1] = f"{key} = {replacements[key]}"
        cell.source = "\n".join(lines)
    return notebook


def function_environment(base):
    notebook = modified_notebook(base)
    env = {"__name__": "rotation_tests"}
    import sys
    import types
    module = types.ModuleType(env["__name__"])
    sys.modules[env["__name__"]] = module
    env = module.__dict__
    for cell in notebook.cells:
        if cell.cell_type != "code":
            continue
        tag = cell.metadata["tags"][0]
        if tag in ["rotation-00", "rotation-01"] or tag.startswith("rotation-core-"):
            exec(cell.source, env)
        elif tag in ["rotation-07", "rotation-08", "rotation-09"]:
            tree = ast.parse(cell.source)
            body = []
            for node in tree.body:
                if tag == "rotation-07" and not isinstance(node, ast.Expr):
                    body.append(node)
                elif isinstance(node, ast.FunctionDef):
                    body.append(node)
            exec(compile(ast.Module(body=body, type_ignores=[]), tag, "exec"), env)
    return env


class RotationTests(unittest.TestCase):
    def test_rotation_model_spaces_and_native_solver_parity(self):
        with tempfile.TemporaryDirectory(prefix="cpe-rotation-unit-") as temp:
            base = Path(temp)
            fixture_files(base)
            env = function_environment(base)
            mapping = env["term_mapping"]
            x_map = mapping.loc[mapping["case"].eq("X21_NATIVE")]
            y_map = mapping.loc[mapping["case"].eq("ROT_Y17_NATIVE")]
            self.assertEqual(x_map["shared_monomial"].sum(), 12)
            self.assertEqual(y_map["shared_monomial"].sum(), 12)
            self.assertFalse(((x_map.power_x == 4) & (x_map.power_y == 0)).any())
            self.assertTrue(((y_map.power_x == 4) & (y_map.power_y == 0)).any())
            data = env["normalize_keys"](fixture_frame())
            data = data.loc[data.Wafer.isin(["1", "2", "3"])].reset_index(drop=True)
            data["rotation_valid_x"] = np.isfinite(data["raw_minus_total_fit_x"])
            data["source_row_id"] = np.arange(len(data))
            untouched = data.copy(deep=True)
            points, coefficients, diagnostics = env["fit_rotation_mode"](data, "OCO")
            pd.testing.assert_frame_equal(data, untouched)
            self.assertTrue(diagnostics.full_column_rank.all())
            for wafer, better, worse in [("1", "ROT_Y17_OLS_DIAGNOSTIC", "X21_OLS_DIAGNOSTIC"), ("2", "X21_OLS_DIAGNOSTIC", "ROT_Y17_OLS_DIAGNOSTIC")]:
                subset = points.loc[points.Wafer.eq(wafer) & points.rotation_valid_x]
                self.assertLess(np.std(subset[f"{better}_residual_x"]), 1e-12)
                self.assertGreater(np.std(subset[f"{worse}_residual_x"]), 1e-5)
            self.assertTrue(points.loc[~points.rotation_valid_x, "ROT_Y17_NATIVE_fit_x"].isna().all())
            # Independently evaluate every saved mapping as a polynomial in original physical coordinates.
            rx, ry = np.array([-9300., 1700., 11000.]), np.array([16000., -4100., 7300.])
            original, _ = env["cpe_38para"](rx, ry)
            u, v = env["rotate_local"](rx, ry)
            _, rotated = env["cpe_38para"](u, v)
            for case, matrix in [("X21_NATIVE", original), ("ROT_Y17_NATIVE", env["ROTATION_SIGN"]*rotated)]:
                for i, row in enumerate(mapping.loc[mapping["case"].eq(case)].itertuples()):
                    expected = row.sign_to_original_x * rx**row.power_x * ry**row.power_y / row.scale
                    np.testing.assert_allclose(matrix[:,i], expected, rtol=1e-13, atol=1e-16)
            # Compare with the full original two-axis shot routine, with Y=0 and identical valid rows/settings.
            source = json.loads((HERE.parent / "MRC_PSM_decorrect.ipynb").read_text(encoding="utf-8"))
            for cell in source["cells"]:
                if cell["cell_type"] != "code":
                    continue
                tree = ast.parse(cell["source"])
                for node in tree.body:
                    if isinstance(node, ast.FunctionDef) and node.name == "run_cpe38_vba_style_shotwise":
                        exec(compile(ast.Module(body=[node], type_ignores=[]), "original_shot_solver", "exec"), env)
            baseline = data.copy()
            baseline["raw_minus_total_fit_y"] = 0.
            actual, _, _ = env["run_cpe38_vba_style_shotwise"](baseline, gopt_spec=.8, gopt_type=1)
            np.testing.assert_allclose(points.X21_NATIVE_fit_x, actual.CPE_fit_x, equal_nan=True, rtol=1e-10, atol=1e-12)
            # Opposite 90-degree sign, including target and inverse transform, gives the same physical X field.
            env["ROTATION_DEG"], env["ROTATION_SIGN"] = -90, -1
            env["term_mapping"] = env["model_table"]()
            reverse, _, _ = env["fit_rotation_mode"](data, "OCO")
            np.testing.assert_allclose(points.ROT_Y17_NATIVE_fit_x, reverse.ROT_Y17_NATIVE_fit_x, equal_nan=True, rtol=1e-9, atol=1e-11)
            clipped = diagnostics.loc[diagnostics.fcp_x.eq(140000.)]
            self.assertTrue((clipped.n_virtual < clipped.n_virtual_full).all())
            self.assertTrue(clipped.groupby(env["SHOT_KEYS"]).n_virtual.nunique().eq(1).all())

    def test_real_kernel_file_filters_statistics_and_excel(self):
        with tempfile.TemporaryDirectory(prefix="cpe-rotation-kernel-") as temp:
            base = Path(temp)
            fixture_files(base)
            notebook = modified_notebook(base)
            notebook.cells.append(nbformat.v4.new_code_cell('''
assert G_OPT_SPEC == .8
assert "raw_minus_total_fit_y" not in frames["ADI"]
assert set(selection["Wafer"]) == {"1", "2", "3"}
assert set(selection["photo_transn_seq"]) == {"9007199254740993"}
assert (diagnostic_table["max_decomposition_error"] < 1e-12).all()
assert wafer_statistics["variance_budget_error_nm2"].abs().max() < 1e-10
for mode, points in results.items():
    for row in wafer_statistics.loc[wafer_statistics["mode"].eq(mode)].itertuples():
        wafer = points.loc[points.Wafer.eq(row.Wafer)]
        values = wafer[f"{row.case}_residual_x"].dropna().to_numpy()*1000
        assert np.isclose(row.residual_three_sigma_nm, 3*np.std(values, ddof=1))
print("Rotation integration checks passed")
'''))
            with patch.dict(os.environ, {"IPYTHONDIR": str(base / "ipython"), "JUPYTER_RUNTIME_DIR": str(base / "jupyter-runtime")}):
                executed = NotebookClient(notebook, kernel_name="python3", timeout=180).execute(cwd=str(base))
            out = base / "output_cpe_critical_rotation"
            points = pd.read_parquet(out / "tables/08_OCO_ROTATION_POINTS.parquet")
            self.assertEqual(len(points), 3*2*81+1)
            self.assertFalse(points.Wafer.eq("99").any())
            self.assertEqual(points.source_row_id.nunique(), len(points))
            diagnostics = pd.read_parquet(out / "tables/08_ALL_SHOT_DIAGNOSTICS.parquet")
            self.assertEqual(len(diagnostics), 2*3*2*4)
            manifest = pd.read_csv(out / "input_manifest.csv")
            adi_input = manifest.loc[manifest.path.str.contains("21_ADI")].iloc[0]
            self.assertTrue(adi_input.path.endswith(".xlsx"))
            self.assertIn("arrowinvalid", adi_input.fallback_errors.lower())
            book = load_workbook(out / "CPE_Critical_Rotation_Comparison.xlsx", read_only=False)
            self.assertIn("Term_mapping", book.sheetnames)
            self.assertEqual(book["Wafer_comparison"].max_row, 7)
            self.assertEqual(book["Wafer_comparison"].freeze_panes, "E2")
            for row in book["Wafer_comparison"].iter_rows(min_row=2):
                photo = row[3]
                self.assertEqual(photo.value, BIG_PHOTO)
                self.assertEqual(photo.data_type, "s")
            book.close()
            self.assertEqual(len(list((out / "plots").glob("*.png"))), 4)
            artifact = Path("/tmp/cpe-rotation-test")
            artifact.mkdir(exist_ok=True)
            nbformat.write(executed, artifact / "executed.ipynb")
            for path in (out / "plots").glob("*.png"):
                shutil.copy2(path, artifact / path.name)
            shutil.copy2(out / "CPE_Critical_Rotation_Comparison.xlsx", artifact / "synthetic_comparison.xlsx")

    def test_input_safety_and_rank_diagnostics(self):
        with tempfile.TemporaryDirectory(prefix="cpe-rotation-errors-") as temp:
            base = Path(temp)
            fixture_files(base)
            env = function_environment(base)
            with self.assertRaises(ValueError):
                env["apply_filters"](pd.DataFrame({"LOT_ID": ["A"]}), "OCO")
            sample = env["STAGE_DIR"] / "SAMPLE_ONLY_sample_1.xlsx"
            pd.DataFrame({"x": [1]}).to_excel(sample, index=False)
            with self.assertRaises(FileNotFoundError):
                env["read_full"](["SAMPLE_ONLY"])
            few = env["normalize_keys"](fixture_frame().head(3))
            few["rotation_valid_x"] = True
            _, _, diagnostics = env["fit_rotation_mode"](few, "OCO")
            self.assertFalse(diagnostics.full_column_rank.any())
            self.assertTrue(np.isinf(diagnostics.column_scaled_condition).all())


if __name__ == "__main__":
    unittest.main()
