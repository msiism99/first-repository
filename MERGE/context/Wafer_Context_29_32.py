# %% 29. Wafer context 설정 / 저장 파일 읽기
# 앞의 DB 조회 셀을 실행하지 않아도 이 셀부터 실행할 수 있습니다.
from pathlib import Path
from decimal import Decimal, InvalidOperation
import json
import re
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from IPython.display import display, FileLink

CTX_BASE_DIR = Path(r"C:\PHOTO_ANALYSIS\BASIC")
CTX_STAGE_DIR = CTX_BASE_DIR / "output_decorrect_notebook" / "stages"
CTX_OUTPUT_DIR = CTX_BASE_DIR / "output_wafer_context"
CTX_KEYS = ["LOT_ID", "Wafer", "photo_transn_seq"]
CTX_EXTRA_COLUMNS = []  # 자동 제외되는 컬럼도 필요하면 원본 이름을 추가하세요.
CTX_CELL_LIMIT = 30000
CTX_ALIASES = {"LOT_ID": ["LOT_ID", "lotid", "lot_id"], "Wafer": ["Wafer", "slotid", "slot_id", "wafer"], "photo_transn_seq": ["photo_transn_seq"]}
CTX_LINK_IDS = ["lot_transn_seq", "apc_hist_index_no", "apc_trocs_hist_index_no"]
CTX_PROFILE = ["fcp_x", "fcp_y", "TEST", "Is_Edge_Shot"]
CTX_POINT_COLUMNS = set(CTX_PROFILE + ["DieX", "DieY", "coordinate_x", "coordinate_y", "coordinate_X", "coordinate_Y", "Coordinate_X", "Coordinate_Y", "value_x", "value_y", "mrc_x_valn", "mrc_y_valn", "wf_x", "wf_y", "radius", "Max_Corner_Distance", "UNIQUE_ID2", "UNIQUE_ID4", "pos_x_valn", "pos_y_valn", "fcp_x_round", "fcp_y_round", "match_key", "trocs_key", "rn"])
ctx_cache, ctx_profiles, ctx_batches_saved, ctx_files, ctx_joins, ctx_columns, ctx_values, ctx_multi = {}, {}, {}, [], [], [], [], []
if not CTX_STAGE_DIR.is_dir():
    raise FileNotFoundError(f"stages 폴더 없음: {CTX_STAGE_DIR}")
CTX_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

def ctx_id(value):
    if pd.isna(value) or str(value).strip() in ["", "<NA>", "None", "nan", "NaT"]:
        return pd.NA
    text = str(value).strip()
    try:
        number = Decimal(text)
        return str(int(number)) if number.is_finite() and number == number.to_integral_value() else text
    except (InvalidOperation, ValueError, OverflowError):
        return text

def ctx_normalize(frame):
    out = frame.copy()
    for key, aliases in CTX_ALIASES.items():
        existing = [col for col in aliases if col in out]
        if not existing:
            continue
        values = [ctx_ids(out[col]) if key != "LOT_ID" else out[col].astype("string").str.strip().replace("", pd.NA) for col in existing]
        merged = values[0]
        for value in values[1:]:
            conflict = merged.notna() & value.notna() & merged.ne(value).fillna(False)
            if conflict.any():
                raise ValueError(f"{key}: 키 별칭 값 충돌 {int(conflict.sum())}행")
            merged = merged.fillna(value)
        out[key] = merged.astype("string")
    for col in CTX_LINK_IDS:
        if col in out:
            out[col] = ctx_ids(out[col])
    return out

def ctx_ids(series):
    text = series.astype("string").str.strip()
    lookup = {value: ctx_id(value) for value in text.dropna().unique()}
    return text.map(lookup).astype("string")

def ctx_is_context(col):
    if col in CTX_EXTRA_COLUMNS:
        return True
    if col in CTX_POINT_COLUMNS or col.startswith("_"):
        return False
    lower = col.lower()
    excluded = r"^(?:x_reg|y_reg|raw_[xy]|raw_minus_|real_residual_|mrc_fit_|psm_fit_[xy]|psm_[xy]$|mrc_r[xy]$|wk_fit_|rk_fit_|total_fit_|cpe_fit_|trocs_fit_|rms_|span_|lambda_|wk\d|rk\d|mrc_k\d|reg_|sim_reg_|oco_sim_)"
    return re.match(excluded, lower) is None

def ctx_read(stem, required=False):
    if stem in ctx_cache:
        return ctx_cache[stem]
    candidates = [CTX_STAGE_DIR / f"{stem}.parquet", CTX_STAGE_DIR / f"{stem}.xlsx"]
    available = [path for path in candidates if path.is_file()]
    samples = list(CTX_STAGE_DIR.glob(f"{stem}_sample_*.xlsx"))
    failures = []
    for path in available:
        try:
            if path.suffix == ".parquet":
                names = pq.read_schema(path).names
            else:
                book = load_workbook(path, read_only=True, data_only=True)
                try:
                    names = list(next(book.active.iter_rows(values_only=True)))
                finally:
                    book.close()
            selected = [col for col in names if isinstance(col, str) and (ctx_is_context(col) or col in CTX_PROFILE)]
            pieces, profile, batches, n_rows, invalid_keys = [], {}, {}, 0, 0
            for chunk in ctx_batches(path, selected):
                chunk = ctx_normalize(chunk)
                n_rows += len(chunk)
                if all(col in chunk for col in CTX_KEYS):
                    invalid_keys += int(chunk[CTX_KEYS].isna().any(axis=1).sum())
                if all(col in chunk for col in CTX_KEYS + ["TEST", "fcp_x", "fcp_y"]):
                    ctx_profile_update(profile, chunk)
                if all(col in chunk for col in ["lot_transn_seq", "Wafer", "data_kind", "transn_occur_date"]):
                    ctx_batch_update(batches, chunk)
                pieces.append(chunk.drop(columns=[col for col in CTX_PROFILE if col not in CTX_EXTRA_COLUMNS], errors="ignore").drop_duplicates())
            frame = pd.concat(pieces, ignore_index=True).drop_duplicates() if pieces else ctx_normalize(pd.DataFrame(columns=selected))
            ctx_profiles[stem] = profile
            ctx_batches_saved[stem] = batches
            ctx_files.append(dict(source=stem, path=str(path), status="read", rows=n_rows, context_rows=len(frame), invalid_wafer_keys=invalid_keys, columns_read=len(selected), bytes=path.stat().st_size, mtime_ns=path.stat().st_mtime_ns, detail="; ".join(failures)))
            ctx_cache[stem] = frame
            return frame
        except Exception as exc:
            failures.append(f"{path.name}: {type(exc).__name__}: {exc}")
    status = "read_error" if failures else "sample_only" if samples else "missing"
    ctx_files.append(dict(source=stem, status=status, detail="; ".join(failures) or f"전체 파일 없음; sample {len(samples)}개는 제외"))
    if required:
        raise FileNotFoundError(f"{stem}: 전체 Parquet/Excel 필요. {status}: {'; '.join(failures)}")
    ctx_cache[stem] = None
    return None

# %% 30. Wafer별 context 집계 / 노광 이력·APC ID 연결 함수
def ctx_batches(path, columns):
    if path.suffix == ".parquet":
        for batch in pq.ParquetFile(path).iter_batches(batch_size=100000, columns=columns):
            yield batch.to_pandas(types_mapper=pd.ArrowDtype)
    else:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = book.active.iter_rows(values_only=True)
            header = list(next(rows))
            indices, buffer = [header.index(col) for col in columns], []
            for row in rows:
                buffer.append([row[index] for index in indices])
                if len(buffer) >= 100000:
                    yield pd.DataFrame(buffer, columns=columns, dtype=object)
                    buffer = []
            if buffer:
                yield pd.DataFrame(buffer, columns=columns, dtype=object)
        finally:
            book.close()

def ctx_profile_update(profile, frame):
    for key, group in frame.dropna(subset=CTX_KEYS).groupby(CTX_KEYS, dropna=False, sort=False):
        item = profile.setdefault(key, dict(n_rows=0, shots=set(), tests=set(), edge_shots=set()))
        item["n_rows"] += len(group)
        item["shots"].update(map(tuple, group[["fcp_x", "fcp_y"]].dropna().drop_duplicates().to_numpy()))
        item["tests"].update(group["TEST"].dropna().unique())
        if "Is_Edge_Shot" in group:
            edge = group["Is_Edge_Shot"].astype("string").str.lower().isin(["true", "1"])
            item["edge_shots"].update(map(tuple, group.loc[edge, ["fcp_x", "fcp_y"]].dropna().drop_duplicates().to_numpy()))

def ctx_batch_update(batches, frame):
    for key, group in frame.groupby(["lot_transn_seq", "Wafer", "data_kind"], dropna=True, sort=False):
        dates = pd.to_datetime(group["transn_occur_date"], errors="coerce").dropna()
        item = batches.setdefault(key, dict(start_time=pd.NaT, end_time=pd.NaT, n_rows=0))
        item["n_rows"] += len(group)
        if len(dates):
            item["start_time"] = dates.min() if pd.isna(item["start_time"]) else min(item["start_time"], dates.min())
            item["end_time"] = dates.max() if pd.isna(item["end_time"]) else max(item["end_time"], dates.max())

def ctx_text(value):
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if pd.isna(value):
        return None
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat(sep=" ")
    return str(value).strip() or None

def ctx_context_columns(frame, keys):
    aliases = {col for values in CTX_ALIASES.values() for col in values}
    return [col for col in frame if col not in keys and col not in aliases and ctx_is_context(col)]

def ctx_aggregate(frame, prefix, source):
    columns = ctx_context_columns(frame, CTX_KEYS)
    work = frame[CTX_KEYS + columns].copy()
    for col in columns:
        work[col] = work[col].map(ctx_text)
    work = work.drop_duplicates()
    records = []
    for key, group in work.groupby(CTX_KEYS, dropna=False, sort=False):
        row = dict(zip(CTX_KEYS, key))
        for col in columns:
            values = sorted(group[col].dropna().unique().tolist())
            name = f"{prefix}__{col}"
            text = " | ".join(values) if values else pd.NA
            overflow = isinstance(text, str) and len(text) > CTX_CELL_LIMIT
            row[name] = f"[{len(values)} values; see Context_values]" if overflow else text
            for value in values:
                ctx_values.append(dict(zip(CTX_KEYS, key), column=name, source=source, source_column=col, value=value))
            if len(values) > 1:
                ctx_multi.append(dict(zip(CTX_KEYS, key), column=name, n_values=len(values), cell_overflow=overflow))
        records.append(row)
    for col in columns:
        ctx_columns.append(dict(column=f"{prefix}__{col}", source=source, source_column=col, dtype=str(frame[col].dtype), rule="유일값 보존; 여러 값은 |로 병기, 원값은 Context_values에 보존"))
    return pd.DataFrame(records, columns=CTX_KEYS + [f"{prefix}__{col}" for col in columns])

def ctx_attach(master, frame, links, join_keys, prefix, source):
    if frame is None or frame.empty:
        ctx_joins.append(dict(source=source, prefix=prefix, join_keys=" + ".join(join_keys), matched_wafers=0, status="no_rows"))
        return master
    if any(col not in frame or col not in links for col in join_keys):
        ctx_joins.append(dict(source=source, prefix=prefix, join_keys=" + ".join(join_keys), matched_wafers=0, status="missing_join_columns"))
        return master
    context_cols = ctx_context_columns(frame, join_keys)
    source_keys = list(dict.fromkeys(join_keys + context_cols))
    rhs = frame[source_keys].dropna(subset=join_keys).drop_duplicates()
    link_cols = list(dict.fromkeys(CTX_KEYS + join_keys))
    lhs = links[link_cols].dropna(subset=link_cols).drop_duplicates()
    # LOT/slot만 있는 측정 batch는 노광 이력이 유일한 경우에만 연결합니다.
    ambiguous = lhs.groupby(join_keys, dropna=False).size()
    n_ambiguous = 0
    if "lot_transn_seq" in join_keys and "photo_transn_seq" not in join_keys:
        bad = ambiguous[ambiguous > 1].reset_index()[join_keys]
        n_ambiguous = len(bad)
        if not bad.empty:
            lhs = lhs.merge(bad.assign(_ambiguous=True), on=join_keys, how="left")
            lhs = lhs.loc[lhs["_ambiguous"].isna()].drop(columns="_ambiguous")
    matched = lhs.merge(rhs, on=join_keys, how="inner", validate="m:m")
    aggregated = ctx_aggregate(matched, prefix, source)
    ctx_joins.append(dict(source=source, prefix=prefix, join_keys=" + ".join(join_keys), matched_wafers=len(aggregated), ambiguous_batch_keys=n_ambiguous, status="matched" if len(aggregated) else "unmatched"))
    return master.merge(aggregated, on=CTX_KEYS, how="left", validate="1:1")

def ctx_point_profile(stem, mode):
    rows = [dict(zip(CTX_KEYS, key), **{f"{mode}__n_point_rows": item["n_rows"], f"{mode}__n_shots": len(item["shots"]), f"{mode}__n_tests": len(item["tests"]), f"{mode}__n_edge_shots": len(item["edge_shots"])}) for key, item in ctx_profiles.get(stem, {}).items()]
    return pd.DataFrame(rows) if rows else pd.DataFrame(columns=CTX_KEYS)

def ctx_choose_anchor(mode):
    choices = [f"15_{mode}_FINAL", f"27_{mode}_FINAL_POINTS", f"19_{mode}_TROCS_ATTACHED", f"14_{mode}_RAW_BEFORE_FILTER", f"08_{mode}_MEASURED", f"04_LOTINFO_{mode}_SELECTED"]
    for stem in choices:
        frame = ctx_read(stem)
        if frame is not None and not frame.empty and all(col in frame for col in CTX_KEYS):
            valid = frame.dropna(subset=CTX_KEYS).copy()
            if not valid.empty:
                return valid, stem
    raise FileNotFoundError(f"{mode}: wafer 식별키가 있는 전체 stage 파일이 없습니다.")

# %% 31. 저장된 table에서 wafer context 통합
ctx_cache, ctx_profiles, ctx_batches_saved, ctx_files = {}, {}, {}, []
ctx_joins, ctx_columns, ctx_values, ctx_multi = [], [], [], []
ctx_anchors, ctx_anchor_names, ctx_links, ctx_lotinfo = {}, {}, {}, {}
for mode in ["ADI", "OCO"]:
    anchor, anchor_name = ctx_choose_anchor(mode)
    ctx_anchors[mode], ctx_anchor_names[mode] = anchor, anchor_name
    ctx_lotinfo[mode] = ctx_read(f"04_LOTINFO_{mode}_SELECTED")
    link = anchor[[col for col in CTX_KEYS + CTX_LINK_IDS if col in anchor]].drop_duplicates()
    lot = ctx_lotinfo[mode]
    if lot is not None and "lot_transn_seq" in lot and "lot_transn_seq" not in link:
        mapping = lot[CTX_KEYS + ["lot_transn_seq"]].drop_duplicates()
        link = link.merge(mapping, on=CTX_KEYS, how="left", validate="m:m")
    ctx_links[mode] = link

wafer_context = pd.concat([frame[CTX_KEYS] for frame in ctx_anchors.values()], ignore_index=True).drop_duplicates().sort_values(CTX_KEYS).reset_index(drop=True)
for mode in ["ADI", "OCO"]:
    anchor, stem, links = ctx_anchors[mode], ctx_anchor_names[mode], ctx_links[mode]
    present = anchor[CTX_KEYS].drop_duplicates().assign(**{f"{mode}_present": True})
    wafer_context = wafer_context.merge(present, on=CTX_KEYS, how="left", validate="1:1")
    wafer_context[f"{mode}_present"] = wafer_context[f"{mode}_present"].eq(True)
    wafer_context = ctx_attach(wafer_context, anchor, anchor[CTX_KEYS], CTX_KEYS, mode, stem)
    wafer_context = wafer_context.merge(ctx_point_profile(stem, mode), on=CTX_KEYS, how="left", validate="1:1")
    lot_stem = f"04_LOTINFO_{mode}_SELECTED"
    wafer_context = ctx_attach(wafer_context, ctx_lotinfo[mode], anchor[CTX_KEYS], CTX_KEYS, f"{mode}_LOTINFO", lot_stem)
    for extra_stem, prefix in [("04_LOTINFO_PAIRED", "PAIR"), ("07_EXPO_SELECTED", "EXPO"), ("07_F1_QUERY", "F1"), (f"09_MRC_{mode}_INPUT", "MRC")]:
        extra = ctx_read(extra_stem)
        if extra_stem == "07_F1_QUERY" and extra is not None and all(col in extra for col in CTX_KEYS + ["impala_insert_time"]):
            extra = extra.sort_values("impala_insert_time").drop_duplicates(CTX_KEYS, keep="last")
        wafer_context = ctx_attach(wafer_context, extra, anchor[CTX_KEYS], CTX_KEYS, f"{mode}_{prefix}", extra_stem)
    for extra_stem, prefix, history in [("11_APC_MATCHED", "APC", "apc_hist_index_no"), ("11_APC_MAIN_QUERY", "APC_MAIN", "apc_hist_index_no"), ("11_APC_EEXY_QUERY", "APC_EEXY", "apc_trocs_hist_index_no"), ("12_PSM_INPUT", "PSM", "apc_hist_index_no"), ("16_TROCS_INPUT_UM", "TROCS", "apc_trocs_hist_index_no")]:
        wafer_context = ctx_attach(wafer_context, ctx_read(extra_stem), links, ["LOT_ID", history], f"{mode}_{prefix}", extra_stem)
    batch_stem = f"06_RAWDATA_{mode}_LATEST"
    batch = ctx_read(batch_stem)
    batch_rows = [dict(zip(["lot_transn_seq", "Wafer", "data_kind"], key), **item) for key, item in ctx_batches_saved.get(batch_stem, {}).items()]
    if batch_rows:
        for kind, summary in pd.DataFrame(batch_rows).groupby("data_kind", dropna=False):
            wafer_context = ctx_attach(wafer_context, summary.drop(columns="data_kind"), links, ["lot_transn_seq", "Wafer"], f"{mode}_BATCH_{kind}", batch_stem)

for col in wafer_context:
    if col not in {item["column"] for item in ctx_columns}:
        ctx_columns.append(dict(column=col, source="derived", source_column=col, dtype=str(wafer_context[col].dtype), rule="식별키 / 방향별 존재 여부 / 원본 실측 행·shot·TEST 개수 집계"))
context_columns = pd.DataFrame(ctx_columns).drop_duplicates("column")
context_values = pd.DataFrame(ctx_values, columns=CTX_KEYS + ["column", "source", "source_column", "value"]).drop_duplicates()
context_multiple = pd.DataFrame(ctx_multi, columns=CTX_KEYS + ["column", "n_values", "cell_overflow"])
context_sources, context_joins = pd.DataFrame(ctx_files), pd.DataFrame(ctx_joins)
wafer_context["context_has_multiple_values"] = wafer_context[CTX_KEYS].merge(context_multiple[CTX_KEYS].drop_duplicates().assign(multiple=True), on=CTX_KEYS, how="left")["multiple"].eq(True).to_numpy()
context_columns = pd.concat([context_columns, pd.DataFrame([dict(column="context_has_multiple_values", source="derived", rule="여러 값을 가진 context 항목이 하나라도 있는지")])], ignore_index=True)
for table in [wafer_context, context_values, context_multiple]:
    for col in ["Wafer", "photo_transn_seq"]:
        try:
            table[col] = pd.array([int(value) for value in table[col]], dtype="Int64")
        except (ValueError, TypeError, OverflowError):
            pass
ctx_front = CTX_KEYS + ["ADI_present", "OCO_present"]
for field in ["P_EQPID", "ReticleID", "Photo_PPID", "ChuckID", "Base_EQP1", "MMO_MRC_EQP", "STEPSEQ", "M_STEPSEQ", "P_TIME", "M_TIME", "GROUP"]:
    ctx_front.extend(f"{mode}__{field}" for mode in ["ADI", "OCO"] if f"{mode}__{field}" in wafer_context)
wafer_context = wafer_context[ctx_front + [col for col in wafer_context if col not in ctx_front]].sort_values(CTX_KEYS).reset_index(drop=True)
context_columns = context_columns.set_index("column").loc[wafer_context.columns].rename_axis("column").reset_index()
assert not wafer_context.duplicated(CTX_KEYS).any(), "wafer context의 식별키가 중복되었습니다."
print(f"Wafer context: {len(wafer_context):,} rows × {len(wafer_context.columns):,} columns")
display(wafer_context.head(10))
display(context_joins)

# %% 32. Excel / 전체 Parquet·CSV 저장
context_notes = pd.DataFrame([
    ["행 기준", "LOT_ID + Wafer + photo_transn_seq. 노광 이력이 다르면 같은 LOT/wafer라도 별도 행입니다."],
    ["대상", "ADI/OCO anchor 파일에 있는 wafer의 합집합. ADI_present/OCO_present로 존재 여부를 확인합니다."],
    ["출처", "ADI__/OCO__는 대표 point 파일, LOTINFO/EXPO/F1/MRC/APC/PSM/TROCS는 해당 저장 table입니다."],
    ["연결", "APC/PSM/TROCS는 LOT와 저장된 이력 ID로 연결합니다. LOT만으로 연결하지 않습니다."],
    ["여러 값", "같은 wafer의 값이 여러 개면 |로 병기합니다. 임의의 첫 값을 고르지 않습니다. Multiple_values와 Context_values에서 확인하세요."],
    ["제외", "측정값, point 좌표, K 계수는 context에서 제외합니다. 추가하려면 CTX_EXTRA_COLUMNS에 원본 이름을 지정하세요."],
    ["미저장 정보", "저장 파일에 없는 device/product 등의 정보는 새로 조회하거나 추정하지 않습니다."],
    ["입력 파일", "stages의 전체 Parquet → 전체 Excel 순서로 읽습니다. sample Excel은 제외합니다."],
    ["데이터 기준", "ADI 실제 오차=raw, OCO 비교 기준=Reg. 이 단계는 측정값이나 보정값을 변경하지 않습니다."],
], columns=["항목", "설명"])
for name, frame in [("wafer_context", wafer_context), ("wafer_context_values", context_values)]:
    frame.to_parquet(CTX_OUTPUT_DIR / f"{name}.parquet", index=False)
    frame.to_csv(CTX_OUTPUT_DIR / f"{name}.csv", index=False, encoding="utf-8-sig")
context_sources.to_csv(CTX_OUTPUT_DIR / "context_sources.csv", index=False, encoding="utf-8-sig")
context_joins.to_csv(CTX_OUTPUT_DIR / "context_joins.csv", index=False, encoding="utf-8-sig")
context_xlsx = CTX_OUTPUT_DIR / "wafer_context.xlsx"
context_sheets = {"Wafer_context": wafer_context, "Columns": context_columns, "Sources": context_sources, "Joins": context_joins, "Multiple_values": context_multiple, "Context_values": context_values, "Notes": context_notes}
with pd.ExcelWriter(context_xlsx, engine="openpyxl") as writer:
    for sheet_name, frame in context_sheets.items():
        if len(frame) > 1048575 or len(frame.columns) > 16384:
            raise ValueError(f"{sheet_name}: Excel 한도를 초과합니다. 이미 저장한 전체 Parquet/CSV를 사용하세요.")
        if frame.select_dtypes(include=["object", "string"]).map(lambda value: isinstance(value, str) and len(value) > 32767).any().any():
            raise ValueError(f"{sheet_name}: Excel cell 한도 초과. 전체 Parquet/CSV에는 원값이 저장되어 있습니다.")
        excel_frame = frame.copy()
        if "photo_transn_seq" in excel_frame:
            excel_frame["photo_transn_seq"] = excel_frame["photo_transn_seq"].astype("string")
        excel_frame.to_excel(writer, sheet_name=sheet_name, index=False)
        sheet = writer.sheets[sheet_name]
        sheet.freeze_panes = "D2" if sheet_name == "Wafer_context" else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="24476A")
            sheet.column_dimensions[cell.column_letter].width = min(42, max(16, len(str(cell.value))+2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top")
(CTX_OUTPUT_DIR / "context_settings.json").write_text(json.dumps(dict(keys=CTX_KEYS, stages=str(CTX_STAGE_DIR), anchors=ctx_anchor_names, extra_columns=CTX_EXTRA_COLUMNS, sources=ctx_files), ensure_ascii=False, indent=2), encoding="utf-8")
print("저장 완료:", context_xlsx)
display(FileLink(str(context_xlsx)))
