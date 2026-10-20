"""Real-kernel metadata integration tests using saved synthetic stages only."""
import ast
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import nbformat
from nbclient import NotebookClient
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
BIG_PHOTO = 9007199254740993


def fixtures(base):
    stage = base / 'output_decorrect_notebook/stages'
    stage.mkdir(parents=True)

    def point(wafer, photo, apc, index, mode):
        return dict(LOT_ID='LOT_A', Wafer=wafer, photo_transn_seq=photo, P_EQPID=mode+'_EQP', ReticleID='RET_A' if index % 2 else 'RET_B', Photo_PPID='RECIPE', ChuckID='CH1', M_TIME=pd.Timestamp('2026-10-16 09:00'), P_TIME=pd.Timestamp('2026-10-16 08:00'), apc_hist_index_no=apc, apc_trocs_hist_index_no=apc+800, GROUP='E1' if mode == 'ADI' else 'E2', TEST=index % 5, fcp_x=float(index % 2), fcp_y=0., Is_Edge_Shot=index % 2 == 0, raw_x=123., X_reg=456., Y_reg=0., WK1_X=777., MRC_RX=4., operator_note='=literal_text', product_family='PRODUCT_X')

    rows = [point(9, 100, 110, i, 'ADI') for i in range(80)]
    for index, row in enumerate(rows):
        row['long_context'] = str(index) + '-' + 'a'*500
    rows.extend(point(9, BIG_PHOTO, 120, i, 'ADI') for i in range(4))
    rows.extend(point(11, 100, 110, i, 'ADI') for i in range(100005))
    rows.append(point(9, None, 110, 1, 'ADI'))
    adi = pd.DataFrame(rows)
    adi['photo_transn_seq'] = pd.array([r['photo_transn_seq'] for r in rows], dtype='Int64')
    adi.to_parquet(stage / '15_ADI_FINAL.parquet', index=False, row_group_size=20000)
    oco_rows = [point(wafer, photo, apc, i, 'OCO') for wafer, photo, apc in [(9,100,110), (9,BIG_PHOTO,120), (11,100,110)] for i in range(3)]
    oco_rows.append(dict(point(10,300,130,1,'OCO'), LOT_ID='LOT_B'))
    oco = pd.DataFrame(oco_rows)
    oco['photo_transn_seq'] = oco['photo_transn_seq'].astype(str)
    oco.to_excel(stage / '15_OCO_FINAL.xlsx', index=False)
    selected = []
    for wafer, photo, apc in [(9,100,110), (9,BIG_PHOTO,120), (11,100,110)]:
        selected.append(dict(lotid='LOT_A', slotid=wafer, photo_transn_seq=photo, lot_transn_seq=9000, peqpid='PHOTO_EQP', reticleid='LOT_RETICLE', psm_mmo_ref_eqp_id='REF_EQP', per_shot_mrc_id='PSM_PATH'))
    lots = pd.DataFrame(selected)
    for mode in ['ADI','OCO']:
        lots.to_parquet(stage / f'04_LOTINFO_{mode}_SELECTED.parquet', index=False)
    f1 = pd.DataFrame([dict(lotid='LOT_A', slotid=9, photo_transn_seq=100, base_eqp_id1=name, impala_insert_time=pd.Timestamp(date)) for name,date in [('OLD_F1','2026-10-15'), ('LATEST_F1','2026-10-16')]])
    f1.to_parquet(stage / '07_F1_QUERY.parquet', index=False)
    apc = pd.DataFrame([dict(lot_id=lot, apc_hist_index_no=apc_id, apc_trocs_hist_index_no=apc_id+800, eqp_id=eqp, sub_eqp_id='EX1', reticle_id=reticle, run_type='NORMAL', event_tsdt='2026-10-16') for lot,apc_id,eqp,reticle in [('LOT_A',110,'APC_CORRECT','APC_RET'),('LOT_A',120,'APC_BIG_PHOTO','APC_RET2'),('LOT_A',999,'WRONG_OTHER_HISTORY','WRONG'),('LOT_OTHER',110,'WRONG_OTHER_LOT','WRONG')]])
    apc.to_parquet(stage / '11_APC_MATCHED.parquet', index=False)
    apc.drop(columns='apc_trocs_hist_index_no').to_parquet(stage / '11_APC_MAIN_QUERY.parquet', index=False)
    apc[['lot_id','apc_trocs_hist_index_no','eqp_id','event_tsdt']].to_parquet(stage / '11_APC_EEXY_QUERY.parquet', index=False)
    psm = apc.copy().assign(pos_x_valn=0., pos_y_valn=0., mrc_k1_valn=5.)
    (stage / '12_PSM_INPUT.parquet').write_bytes(b'corrupt parquet - must fall back to full excel')
    psm.to_excel(stage / '12_PSM_INPUT.xlsx', index=False)
    psm.assign(eqp_id='SAMPLE_MUST_NOT_BE_USED').head(1).to_excel(stage / '12_PSM_INPUT_sample_1.xlsx', index=False)
    pd.DataFrame({'lot_id':['LOT_A'], 'eqp_id':['SAMPLE_ONLY']}).to_excel(stage / '16_TROCS_INPUT_UM_sample_1.xlsx', index=False)
    batch = pd.DataFrame([dict(lot_transn_seq=9000, slot_id=wafer, data_kind='RAWDATA', transn_occur_date=pd.Timestamp('2026-10-16')+pd.Timedelta(seconds=i), value_x=999.) for wafer,count in [(9,5),(11,8)] for i in range(count)])
    batch.to_parquet(stage / '06_RAWDATA_ADI_LATEST.parquet', index=False)
    return stage


class ContextTests(unittest.TestCase):
    def test_actual_kernel_saved_metadata_join_and_excel(self):
        with tempfile.TemporaryDirectory(prefix='wafer-context-') as temp:
            base = Path(temp)
            fixtures(base)
            notebook = nbformat.read(HERE / 'Wafer_Context_29_32.ipynb', as_version=4)
            nbformat.validate(notebook)
            for cell in notebook.cells:
                if 'wafer-context-29' in cell.metadata.get('tags', []):
                    lines = cell.source.splitlines()
                    node = next(n for n in ast.parse(cell.source).body if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'CTX_BASE_DIR' for t in n.targets))
                    lines[node.lineno-1] = f'CTX_BASE_DIR = Path({str(base)!r})'
                    cell.source = '\n'.join(lines)
            notebook.cells.append(nbformat.v4.new_code_cell('''
try:
    ctx_normalize(pd.DataFrame({"LOT_ID":["A"],"lotid":["B"]}))
except ValueError:
    pass
else:
    raise AssertionError("conflicting aliases must not be silently coalesced")
assert ctx_id("9007199254740993.0") == "9007199254740993"
assert set(wafer_context.columns) == set(context_columns["column"])
assert not any("raw_x" in c or "WK1_X" in c or "value_x" in c for c in wafer_context)
print("Context integration assertions passed")
'''))
            build_cell = next(cell.source for cell in notebook.cells if 'wafer-context-31' in cell.metadata.get('tags', []))
            notebook.cells.append(nbformat.v4.new_code_cell('before_rerun = wafer_context.copy(deep=True)\nbefore_multiple = len(context_multiple)\nexec(' + repr(build_cell) + ')\npd.testing.assert_frame_equal(before_rerun, wafer_context)\nassert len(context_multiple) == before_multiple\nprint("Rerunning the build cell is idempotent")'))
            with patch.dict(os.environ, {'IPYTHONDIR':str(base/'ipython'), 'JUPYTER_RUNTIME_DIR':str(base/'jupyter-runtime')}):
                executed = NotebookClient(notebook, kernel_name='python3', timeout=180).execute(cwd=str(base))
            artifact = Path('/tmp/wafer-context-test')
            artifact.mkdir(exist_ok=True)
            nbformat.write(executed, artifact / 'executed.ipynb')
            out = base / 'output_wafer_context'
            table = pd.read_parquet(out / 'wafer_context.parquet')
            self.assertEqual(len(table), 4)
            self.assertFalse(table.duplicated(['LOT_ID','Wafer','photo_transn_seq']).any())
            regular = table.loc[table['LOT_ID'].eq('LOT_A') & table['Wafer'].eq(9) & table['photo_transn_seq'].eq(100)].iloc[0]
            shared = table.loc[table['LOT_ID'].eq('LOT_A') & table['Wafer'].eq(11)].iloc[0]
            big = table.loc[table['photo_transn_seq'].eq(BIG_PHOTO)].iloc[0]
            self.assertEqual(regular['ADI__P_EQPID'], 'ADI_EQP')
            self.assertEqual(regular['OCO__P_EQPID'], 'OCO_EQP')
            self.assertEqual(regular['ADI__ReticleID'], 'RET_A | RET_B')
            self.assertTrue(regular['context_has_multiple_values'])
            self.assertEqual(regular['ADI_F1__base_eqp_id1'], 'LATEST_F1')
            self.assertEqual(regular['ADI_APC__eqp_id'], 'APC_CORRECT')
            self.assertEqual(shared['ADI_APC__eqp_id'], 'APC_CORRECT')
            self.assertEqual(big['ADI_APC__eqp_id'], 'APC_BIG_PHOTO')
            self.assertEqual(regular['ADI_PSM__eqp_id'], 'APC_CORRECT')
            self.assertEqual(regular['ADI__n_point_rows'], 80)
            self.assertEqual(shared['ADI__n_point_rows'], 100005)
            self.assertEqual(shared['ADI__n_shots'], 2)
            self.assertEqual(shared['ADI__n_tests'], 5)
            self.assertEqual(shared['ADI__n_edge_shots'], 1)
            self.assertEqual(shared['ADI_BATCH_RAWDATA__n_rows'], '8')
            self.assertTrue(pd.isna(regular['ADI_BATCH_RAWDATA__n_rows']))
            self.assertTrue(pd.isna(big['ADI_BATCH_RAWDATA__n_rows']))
            self.assertTrue(regular['ADI__long_context'].startswith('[80 values;'))
            only_oco = table.loc[table['LOT_ID'].eq('LOT_B')].iloc[0]
            self.assertFalse(only_oco['ADI_present'])
            self.assertTrue(only_oco['OCO_present'])
            self.assertNotIn('WRONG_OTHER', table.to_csv(index=False))
            sources = pd.read_csv(out / 'context_sources.csv').set_index('source')
            self.assertEqual(sources.loc['16_TROCS_INPUT_UM','status'], 'sample_only')
            self.assertIn('corrupt', sources.loc['12_PSM_INPUT','detail'])
            self.assertEqual(sources.loc['15_ADI_FINAL','invalid_wafer_keys'], 1)
            self.assertLess(sources.loc['15_ADI_FINAL','context_rows'], 100)
            joins = pd.read_csv(out / 'context_joins.csv')
            self.assertEqual(joins.loc[joins['prefix'].eq('ADI_BATCH_RAWDATA'),'ambiguous_batch_keys'].iloc[0], 1)
            values = pd.read_parquet(out / 'wafer_context_values.parquet')
            self.assertEqual(values.loc[values['column'].eq('ADI__long_context') & values['photo_transn_seq'].eq(100)].shape[0], 80)
            book = load_workbook(out / 'wafer_context.xlsx', read_only=False, data_only=False)
            try:
                main = book['Wafer_context']
                self.assertEqual(main.max_row, 5)
                self.assertEqual(main.freeze_panes, 'D2')
                headers = {cell.value:cell.column for cell in main[1]}
                photo_cells = [main.cell(i, headers['photo_transn_seq']) for i in range(2,6)]
                self.assertIn(str(BIG_PHOTO), [cell.value for cell in photo_cells])
                self.assertTrue(all(cell.data_type == 's' for cell in photo_cells))
                literal = main.cell(2, headers['ADI__operator_note'])
                self.assertEqual(literal.value, '=literal_text')
                self.assertEqual(literal.data_type, 's')
                self.assertTrue(main.auto_filter.ref)
            finally:
                book.close()
            (artifact / 'wafer_context.xlsx').write_bytes((out / 'wafer_context.xlsx').read_bytes())
            settings = json.loads((out / 'context_settings.json').read_text())
            self.assertEqual(settings['anchors'], {'ADI':'15_ADI_FINAL','OCO':'15_OCO_FINAL'})


if __name__ == '__main__':
    unittest.main()
