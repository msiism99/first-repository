"""Sampling identifiability and noise transfer tests with actual-kernel integration."""
import ast
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

import nbformat
from nbclient import NotebookClient
import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
PHOTO = "9007199254740993"


def fixtures(base, coordinates_only=False, identical=False):
    stage = base / "output_decorrect_notebook/stages"
    stage.mkdir(parents=True)
    for mode in ["ADI","OCO"]:
        nx, ny = (6,3) if mode == "ADI" or identical else (4,6)
        coordinates = [(x,y) for x in np.linspace(-11000,11000,nx) for y in np.linspace(-14000,14000,ny)]
        coordinates.append(coordinates[len(coordinates)//2])  # repeated measurement, retained as a weight
        rows = []
        for wafer,shot_x,points in [(1,0.,coordinates),(1,20000.,coordinates),(1,140000.,coordinates[:3]),(2,0.,coordinates[:3])]:
            for x,y in points:
                value = .001*(.8*(y/16000.)**3+.6*(x/12000.)**4)
                row = dict(LOT_ID="LOT_S",Wafer=wafer,photo_transn_seq=PHOTO,fcp_x=shot_x,fcp_y=0.,coordinate_X=x,coordinate_Y=y,STEP_PITCH_X=24000.,STEP_PITCH_Y=32000.,CHIP_X_NUM=2,CHIP_Y_NUM=2)
                if not coordinates_only:
                    row["raw_minus_total_fit_x"] = value
                rows.append(row)
        frame = pd.DataFrame(rows)
        if mode == "ADI" or coordinates_only or identical:
            frame.to_parquet(stage / f"21_{mode}_WK_RK_REMOVED.parquet",index=False)
        else:
            (stage / f"21_{mode}_WK_RK_REMOVED.parquet").write_bytes(b"invalid parquet")
            frame.to_excel(stage / f"21_{mode}_WK_RK_REMOVED.xlsx",index=False)
    (stage.parent / "extension_settings.json").write_text(json.dumps(dict(cpe_options=dict(gopt_spec=.9,gopt_type=1),use_virtual_chip_points=True,virtual_wafer_radius_um=144000.)),encoding="utf-8")


def modified_notebook(base):
    notebook = nbformat.read(HERE / "CPE_Sampling_Diagnostics.ipynb",as_version=4)
    nbformat.validate(notebook)
    for cell in notebook.cells:
        if "sampling-00" not in cell.metadata.get("tags",[]):
            continue
        replacements = dict(BASE_DIR=f"Path({str(base)!r})",SHOW_PLOTS="False",COMMON_GRID_N="17",RANDOM_MIXTURES="3")
        lines = cell.source.splitlines()
        for node in ast.parse(cell.source).body:
            if isinstance(node,ast.Assign) and isinstance(node.targets[0],ast.Name) and node.targets[0].id in replacements:
                key = node.targets[0].id
                lines[node.lineno-1] = f"{key} = {replacements[key]}"
        cell.source = "\n".join(lines)
    return notebook


def environment(base, last=9):
    name = "sampling_test_"+str(abs(hash(str(base))))
    module = types.ModuleType(name)
    sys.modules[name] = module
    for cell in modified_notebook(base).cells:
        if cell.cell_type != "code":
            continue
        tag = cell.metadata["tags"][0]
        number = int(tag.rsplit("-",1)[1])
        if number <= last:
            exec(compile(cell.source,tag,"exec"),module.__dict__)
    return module.__dict__


class SamplingTests(unittest.TestCase):
    def test_different_sampling_alias_replay_noise_and_native_parity(self):
        with tempfile.TemporaryDirectory(prefix="sampling-model-") as temp:
            base = Path(temp)
            fixtures(base)
            env = environment(base,last=11)
            summary = env["sampling_summary"].set_index(["mode","case"])
            self.assertEqual(summary.loc[("ADI","X21_NATIVE"),"n_unique_points"],18)
            self.assertEqual(summary.loc[("OCO","X21_NATIVE"),"n_unique_points"],24)
            self.assertEqual(summary.loc[("ADI","X21_NATIVE"),"n_points"],19)
            self.assertLess(summary.loc[("ADI","X21_NATIVE"),"rank"],21)
            self.assertEqual(summary.loc[("OCO","X21_NATIVE"),"rank"],21)
            self.assertEqual(summary.loc[("ADI","ROT_Y17_NATIVE"),"rank"],17)
            self.assertLess(summary.loc[("OCO","ROT_Y17_NATIVE"),"rank"],17)
            self.assertTrue(env["sampling_interaction"].matches_observed_noiseless.any())
            self.assertTrue(env["selection"].common_wafer.all())
            self.assertEqual(set(env["selection"].Wafer),{"1"})
            self.assertEqual(set(env["selection"].photo_transn_seq),{PHOTO})
            self.assertEqual(len(env["layout_catalogue"]),4)
            response = env["pattern_response"]
            for mode,case in [("ADI","ROT_Y17_NATIVE"),("OCO","X21_NATIVE")]:
                pattern = "x4_y0" if mode == "ADI" else "x0_y3"
                row = response.loc[response["mode"].eq(mode) & response.case.eq(case) & response.method.eq("OLS_SVD") & response.pattern.eq(pattern)].iloc[0]
                self.assertLess(row.grid_rmse_nm,1e-10)
            # Native span implementation reuses only the geometry-dependent lambda; match the original solver.
            for (mode,case),spec in env["model_cases"].items():
                work = env["primary_frames"][mode]
                y = work[env["TARGET"]].to_numpy(float)*1000
                beta,_,_ = env["native_fixed_lambda"](spec["matrix"],y,spec["evaluation"],spec["names"],spec["lambdas"])
                direct = env["solve_case"](spec["matrix"],y/1000,spec["evaluation"],spec["names"],len(spec["evaluation"]))
                np.testing.assert_allclose(spec["matrix"] @ beta,direct.pred*1000,rtol=1e-8,atol=1e-9)
                np.testing.assert_allclose(direct.lambdas,spec["lambdas"],rtol=1e-12)
            # Exact linear noise propagation agrees with an independent Monte Carlo realization.
            spec = env["model_cases"][("ADI","ROT_Y17_NATIVE")]
            hat = spec["common_eval"] @ spec["operators"]["GOPT_RIDGE"]
            generator = np.random.default_rng(65)
            noise = generator.normal(size=(hat.shape[1],10000))
            measured_variance = float(np.mean((hat @ noise)**2))
            expected_variance = float(np.mean(np.sum(hat**2,axis=1)))
            self.assertAlmostEqual(measured_variance/expected_variance,1.,delta=.025)
            residual_noise = (np.eye(len(spec["matrix"]))-spec["matrix"] @ spec["operators"]["GOPT_RIDGE"]) @ noise
            sample_variance = np.mean(np.var(residual_noise,axis=0,ddof=1))
            self.assertAlmostEqual(sample_variance/spec["train_noise_gains"]["GOPT_RIDGE"]**2,1.,delta=.025)
            # A rank-deficient sampling can fit its measurements while missing the field between measurements.
            alias_row = response.loc[response["mode"].eq("ADI") & response.case.eq("X21_NATIVE") & response.method.eq("OLS_SVD") & response.pattern.eq("x0_y3")].iloc[0]
            self.assertLess(alias_row.train_rmse_nm,1e-10)
            self.assertGreater(alias_row.grid_rmse_nm,.05)

    def test_identical_sampling_gives_identical_behavior(self):
        with tempfile.TemporaryDirectory(prefix="sampling-same-") as temp:
            base = Path(temp)
            fixtures(base,coordinates_only=True,identical=True)
            env = environment(base,last=11)
            interactions = env["sampling_interaction"]
            self.assertFalse(interactions.matches_observed_noiseless.any())
            self.assertFalse(interactions.matches_observed_with_assumed_noise.any())
            for metric in ["grid_rmse_nm","expected_grid_rmse_nm","shared_hull_rmse_nm"]:
                np.testing.assert_allclose(interactions[f"ADI_delta_{metric}_ROT_minus_X21"],interactions[f"OCO_delta_{metric}_ROT_minus_X21"],rtol=1e-12,atol=1e-12,equal_nan=True)
            self.assertTrue(env["observed_check"].empty)
            for mode in ["ADI","OCO"]:
                self.assertNotIn("raw_minus_total_fit_y",env["primary_frames"][mode])

    def test_real_kernel_coordinates_only_and_excel(self):
        with tempfile.TemporaryDirectory(prefix="sampling-kernel-") as temp:
            base = Path(temp)
            fixtures(base,coordinates_only=True)
            notebook = modified_notebook(base)
            notebook.cells.append(nbformat.v4.new_code_cell('''
assert observed_check.empty
assert observed_group_gain.empty
assert G_OPT_SPEC == .9
assert len(pattern_response) == len(pattern_catalogue)*2*2*3
assert sampling_summary["ridge_noise_gain_rms"].ge(0).all()
assert selection["Wafer"].eq("1").all()
assert all(work["fcp_x"].iloc[0] == 0 for work in primary_frames.values())
assert matching_examples["matches_observed_noiseless"].any()
print("Sampling integration assertions passed")
'''))
            with patch.dict(os.environ,{"IPYTHONDIR":str(base/"ipython"),"JUPYTER_RUNTIME_DIR":str(base/"jupyter-runtime")}):
                executed = NotebookClient(notebook,kernel_name="python3",timeout=180).execute(cwd=str(base))
            out = base / "output_cpe_sampling_diagnostics"
            self.assertEqual(len(list((out/"plots").glob("*.png"))),6)
            self.assertEqual(len(list((out/"plots").glob("*.pdf"))),6)
            book = load_workbook(out/"CPE_Sampling_Diagnostics.xlsx",read_only=False)
            self.assertEqual(book["Selected_wafers"].max_row,3)
            self.assertEqual(book["Sampling_summary"].max_row,5)
            self.assertEqual(book["Observed_check"].max_row,1)
            self.assertEqual(book["Selected_wafers"].cell(2,4).value,PHOTO)
            self.assertEqual(book["Selected_wafers"].cell(2,4).data_type,"s")
            self.assertEqual(book["Sampling_summary"].freeze_panes,"C2")
            book.close()
            artifact = Path("/tmp/cpe-sampling-test")
            artifact.mkdir(exist_ok=True)
            nbformat.write(executed,artifact/"executed.ipynb")
            for path in (out/"plots").glob("*.png"):
                shutil.copy2(path,artifact/path.name)
            shutil.copy2(out/"CPE_Sampling_Diagnostics.xlsx",artifact/"synthetic_sampling.xlsx")
            pd.read_parquet(out/"tables/09_SAMPLING_SUMMARY.parquet").to_csv(artifact/"synthetic_summary.csv",index=False)


if __name__ == "__main__":
    unittest.main()
